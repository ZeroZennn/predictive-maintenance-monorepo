import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'RAB Kainosoph'

# Border
thin = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Fills
header_fill = PatternFill('solid', fgColor='1F4E79')   # Dark blue
cat_fill    = PatternFill('solid', fgColor='2E75B6')   # Medium blue
subcat_fill = PatternFill('solid', fgColor='BDD7EE')   # Light blue
total_fill  = PatternFill('solid', fgColor='FFD966')   # Yellow

# Alignments
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left   = Alignment(horizontal='left',   vertical='center', wrap_text=True)
right  = Alignment(horizontal='right',  vertical='center', wrap_text=True)

# ── Title ────────────────────────────────────────────────────────────────────
ws.merge_cells('A1:G1')
t = ws['A1']
t.value = 'RENCANA ANGGARAN BIAYA (RAB) SUPER DETAIL - PT KAINOSOPH'
t.font  = Font(name='Arial', bold=True, size=13, color='FFFFFF')
t.fill  = header_fill
t.alignment = center
ws.row_dimensions[1].height = 30

ws.merge_cells('A2:G2')
s = ws['A2']
s.value = '(Kebutuhan Investasi Awal / Tahun 0)'
s.font  = Font(name='Arial', italic=True, size=11)
s.alignment = center
ws.row_dimensions[2].height = 20

# ── Column headers ───────────────────────────────────────────────────────────
headers = ['No', 'Kategori & Deskripsi Biaya', 'Kuantitas', 'Satuan',
           'Durasi / Vol', 'Harga Satuan (Rp)', 'Total Biaya (Rp)']
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col, value=h)
    cell.font      = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    cell.fill      = header_fill
    cell.alignment = center
    cell.border    = border
ws.row_dimensions[3].height = 35

# ── Data ─────────────────────────────────────────────────────────────────────
# Each tuple: (No, Deskripsi, Qty, Satuan, Durasi, HargaSatuan, Total, type)
# type: 'cat' | 'subcat' | 'item' | 'total'
data = [
    ('A',   'BIAYA PRA-INVESTASI & AKTIVA TIDAK BERWUJUD',              '',  '',          '',           '',          58500000,    'cat'),
    ('1',   'Jasa Notaris & Akta Pendirian',                            1,   'Paket',     '1 Kali',     10000000,    10000000,    'item'),
    ('2',   'Pengesahan Kemenkumham & BNRI',                            1,   'Paket',     '1 Kali',     2500000,     2500000,     'item'),
    ('3',   'Pendaftaran Merek Kainosoph & PRIME',                      2,   'Sertifikat','1 Kali',     3500000,     7000000,     'item'),
    ('4',   'Hak Cipta Program Komputer (Algoritma)',                   1,   'Sertifikat','1 Kali',     5000000,     5000000,     'item'),
    ('5',   'Sertifikasi PSE Kominfo & Komitmen NIB',                   1,   'Paket',     '1 Kali',     4000000,     4000000,     'item'),
    ('6',   'Pembuatan Draft NDA & B2B SaaS Agreement',                 1,   'Paket',     '1 Kali',     15000000,    15000000,    'item'),
    ('7',   'Jasa Konsultan Hukum (Retainer 3 bln)',                    1,   'Paket',     '3 Bulan',    5000000,     15000000,    'item'),

    ('B',   'BIAYA AKTIVA TETAP BERWUJUD (TANGIBLE ASSETS)',            '',  '',          '',           '',          444300000,   'cat'),
    ('8',   'Peralatan Kerja (Laptop High-End Karyawan)',               14,  'Unit',      '1 Kali',     15000000,    210000000,   'item'),
    ('9',   'Inventaris Kantor (Meja, Kursi, Proyektor, AC, dll)',      1,   'Paket',     '1 Kali',     50000000,    50000000,    'item'),
    ('10',  'Set sensor (Vibrasi & Suhu) model NCD V3',                 10,  'Unit',      '1 Kali',     7300000,     73000000,    'item'),
    ('11',  'NCD Cloud Gateway penerima sinyal',                        1,   'Unit',      '1 Kali',     9500000,     9500000,     'item'),
    ('12',  'Sensor (tekanan) model DCT531i',                           10,  'Unit',      '1 Kali',     7500000,     75000000,    'item'),
    ('13',  'Sensor (rpm) model Optical',                               10,  'Unit',      '1 Kali',     650000,      6500000,     'item'),
    ('14',  'Sensor power consumption model ME237',                     10,  'Unit',      '1 Kali',     380000,      3800000,     'item'),
    ('15',  'Sensor humidity model CWT-TH03S-M',                       10,  'Unit',      '1 Kali',     315000,      3150000,     'item'),
    ('16',  'Sensor noise level model RN-ZS-N01',                      10,  'Unit',      '1 Kali',     1000000,     10000000,    'item'),
    ('17',  'Router Industri',                                          1,   'Unit',      '1 Kali',     3000000,     3000000,     'item'),
    ('18',  'Kabel cat6 (100 meter)',                                   1,   'Unit',      '100 Meter',  3500,        350000,      'item'),

    ('C',   'BIAYA OPERASIONAL (MODAL KERJA 1 TAHUN)',                  '',  '',          '',           '',          2195100000,  'cat'),
    ('C.1', 'Beban Gaji Divisi Teknologi & Produk (12 Bulan)',          '',  '',          '',           '',          '',          'subcat'),
    ('19',  'Project Manager',                                          1,   'Orang',     '12 Bulan',   10000000,    120000000,   'item'),
    ('20',  'Data & ML Engineer',                                       1,   'Orang',     '12 Bulan',   15000000,    180000000,   'item'),
    ('21',  'NLP & GenAI Engineer',                                     1,   'Orang',     '12 Bulan',   15000000,    180000000,   'item'),
    ('22',  'Backend & Data Engineer',                                  1,   'Orang',     '12 Bulan',   12000000,    144000000,   'item'),
    ('23',  'Frontend Engineer',                                        1,   'Orang',     '12 Bulan',   10000000,    120000000,   'item'),
    ('24',  'UI/UX Designer',                                          1,   'Orang',     '12 Bulan',   10000000,    120000000,   'item'),
    ('25',  'QA/QC Specialist',                                        1,   'Orang',     '12 Bulan',   12000000,    144000000,   'item'),
    ('26',  'IoT Integration Engineer',                                 1,   'Orang',     '12 Bulan',   10000000,    120000000,   'item'),

    ('C.2', 'Beban Gaji Divisi Bisnis & Operasional (12 Bulan)',        '',  '',          '',           '',          '',          'subcat'),
    ('27',  'B2B Account Manager & Direct Sales',                       1,   'Orang',     '12 Bulan',   8000000,     96000000,    'item'),
    ('28',  'Digital Marketing Specialist',                             1,   'Orang',     '12 Bulan',   7000000,     84000000,    'item'),
    ('29',  'Project & Implementation Coordinator',                     1,   'Orang',     '12 Bulan',   7000000,     84000000,    'item'),
    ('30',  'Finance & Accounting Officer',                             1,   'Orang',     '12 Bulan',   6000000,     72000000,    'item'),
    ('31',  'Legal & Corporate Secretary',                              1,   'Orang',     '12 Bulan',   6000000,     72000000,    'item'),
    ('32',  'IT Support / Customer Success',                            1,   'Orang',     '12 Bulan',   6000000,     72000000,    'item'),

    ('C.3', 'Infrastruktur Cloud & Server',                             '',  '',          '',           '',          '',          'subcat'),
    ('33',  'AWS EC2 G4dn Instance (LSTM Inference)',                   1,   'Paket',     '12 Bulan',   10000000,    120000000,   'item'),
    ('34',  'Cloud Storage & Managed DB (TimescaleDB)',                  1,   'Paket',     '12 Bulan',   3500000,     42000000,    'item'),
    ('35',  'Domain, SSL, dan Hosting Web',                             1,   'Paket',     '1 Tahun',    1500000,     1500000,     'item'),
    ('36',  'API LLM (GPT-4o)',                                         1,   'Paket',     '12 Bulan',   4000000,     48000000,    'item'),
    ('37',  'Managed Vector Database (Qdrant Cloud)',                   1,   'Paket',     '12 Bulan',   2000000,     24000000,    'item'),

    ('C.4', 'Pemasaran & Akuisisi Klien',                               '',  '',          '',           '',          '',          'subcat'),
    ('38',  'Sewa Lahan Booth Pameran (3x3m) & Listrik',               1,   'Paket',     '1 Kali',     30000000,    30000000,    'item'),
    ('39',  'Konstruksi Booth & Dekorasi Futuristik',                   1,   'Paket',     '1 Kali',     10000000,    10000000,    'item'),
    ('40',  'Marketing Kit (Brosur, Video 3D, Merchandise)',            1,   'Paket',     '1 Kali',     8000000,     8000000,     'item'),
    ('41',  'Iklan Digital B2B (LinkedIn Ads)',                         1,   'Paket',     '12 Bulan',   13500000,    162000000,   'item'),
    ('42',  'SEO Optimization',                                         1,   'Paket',     '12 Bulan',   1500000,     18000000,    'item'),

    ('C.5', 'Perjalanan Dinas & Operasional Kantor',                    '',  '',          '',           '',          '',          'subcat'),
    ('43',  'Biaya Bahan Bakar Kunjungan Pabrik (BBM)',                 10,  'Liter',     '10 Kali',    10000,       1000000,     'item'),
    ('44',  'Uang Lamsam & Makan Tim Kunjungan',                        4,   'Orang',     '1 Hari',     150000,      600000,      'item'),
    ('45',  'Penginapan Tim Instalasi (Area Industri)',                  2,   'Unit',      '1 Hari',     200000,      400000,      'item'),
    ('46',  'Sewa Kantor',                                              1,   'Ruangan',   '1 Tahun',    100000000,   100000000,   'item'),
    ('47',  'Langganan Internet Service',                               1,   'Paket',     '12 Bulan',   1000000,     12000000,    'item'),
    ('48',  'Biaya Listrik dan Air Kantor',                             1,   'Ruangan',   '12 Bulan',   800000,      9600000,     'item'),

    ('',    'TOTAL KEBUTUHAN INVESTASI AWAL (X0)',                      '',  '',          '',           '',          2697900000,  'total'),
]

ROW_START = 4
for idx, item in enumerate(data):
    no, desc, qty, sat, dur, harga, total, rtype = item
    row = ROW_START + idx

    vals = [no, desc, qty if qty != '' else None, sat if sat != '' else None,
            dur if dur != '' else None, harga if harga != '' else None,
            total if total != '' else None]

    for col, val in enumerate(vals, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.border = border

        if rtype == 'cat':
            cell.font      = Font(name='Arial', bold=True, size=11, color='FFFFFF')
            cell.fill      = cat_fill
            cell.alignment = center if col != 2 else left
            ws.row_dimensions[row].height = 30
        elif rtype == 'subcat':
            cell.font      = Font(name='Arial', bold=True, size=11, color='1F4E79')
            cell.fill      = subcat_fill
            cell.alignment = center if col not in [2] else left
            ws.row_dimensions[row].height = 25
        elif rtype == 'total':
            cell.font      = Font(name='Arial', bold=True, size=12)
            cell.fill      = total_fill
            cell.alignment = right if col == 7 else (center if col in [1,3,4,5,6] else left)
            ws.row_dimensions[row].height = 28
        else:
            cell.font      = Font(name='Arial', size=11)
            cell.alignment = right if col in [6, 7] else (center if col in [1, 3, 4, 5] else left)
            ws.row_dimensions[row].height = 22

        # Number format for currency columns
        if col in [6, 7] and isinstance(val, (int, float)):
            cell.number_format = '#,##0'

# ── Column widths ─────────────────────────────────────────────────────────────
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 48
ws.column_dimensions['C'].width = 12
ws.column_dimensions['D'].width = 13
ws.column_dimensions['E'].width = 14
ws.column_dimensions['F'].width = 22
ws.column_dimensions['G'].width = 22

ws.freeze_panes = 'A4'

# ── Save ──────────────────────────────────────────────────────────────────────
out = r'C:\PORTFOLIO\PROJECTS\PBL\Predictive Maintenance\projects\predictive-maintenance-monorepo\kewirus\RAB_Kainosoph.xlsx'
wb.save(out)
print('SAVED:', out)
