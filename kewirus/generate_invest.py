"""
Generate InvestKainosoph.xlsx
Tiga sheet:
  1. Proyeksi Pendapatan   – proyeksi 5 tahun
  2. Payback Period        – kumulatif aliran kas
  3. NPV PI IRR            – present value & metrik kelayakan
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ─── Helpers ────────────────────────────────────────────────────────────────
thin   = Side(style='thin')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def make_fill(hex_color):
    return PatternFill('solid', fgColor=hex_color)

# Colour palette
C_NAVY   = '1F4E79'
C_BLUE   = '2E75B6'
C_LBLUE  = 'BDD7EE'
C_GOLD   = 'FFD966'
C_GREEN  = 'E2EFDA'
C_ORANGE = 'FCE4D6'
C_WHITE  = 'FFFFFF'

def hdr_cell(ws, row, col, value,
             bg=C_NAVY, fg='FFFFFF', bold=True,
             align='center', sz=11):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name='Arial', bold=bold, size=sz, color=fg)
    c.fill      = make_fill(bg)
    c.border    = border
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    return c

def data_cell(ws, row, col, value,
              bg=C_WHITE, bold=False, align='left',
              fmt=None, sz=11):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name='Arial', bold=bold, size=sz)
    c.fill      = make_fill(bg)
    c.border    = border
    c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
    if fmt:
        c.number_format = fmt
    return c

def title_row(ws, row, text, ncols, bg=C_NAVY, fg='FFFFFF', sz=13):
    ws.merge_cells(start_row=row, start_column=1,
                   end_row=row,   end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font      = Font(name='Arial', bold=True, size=sz, color=fg)
    c.fill      = make_fill(bg)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 30

def sub_row(ws, row, text, ncols, bg=C_LBLUE, sz=11):
    ws.merge_cells(start_row=row, start_column=1,
                   end_row=row,   end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font      = Font(name='Arial', italic=True, size=sz, color='1F4E79')
    c.fill      = make_fill(bg)
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 22

NUM_FMT  = '#,##0'        # integer rupiah
PERC_FMT = '0.00%'


# ════════════════════════════════════════════════════════════════════════════
#  SHEET 1 – PROYEKSI PENDAPATAN 5 TAHUN
# ════════════════════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()
ws1 = wb.active
ws1.title = '1. Proyeksi Pendapatan'

# Title
title_row(ws1, 1, 'PROYEKSI PENDAPATAN 5 TAHUN – PT KAINOSOPH (SaaS)', 6)

# Asumsi block
sub_row(ws1, 2, 'Asumsi Dasar:', 6)
assumptions = [
    'Harga Langganan (SaaS): Rp 120.000.000 / mesin / tahun (Rp 10 Juta/bulan)',
    'Modal Awal (X0): Rp 2.697.900.000 – mencakup seluruh kebutuhan modal dan biaya operasional Tahun ke-1',
    'Kenaikan OpEx: +10% per tahun mulai Tahun ke-2 (skalabilitas SDM + kapasitas cloud server)',
]
for i, a in enumerate(assumptions, 3):
    ws1.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    c = ws1.cell(row=i, column=1, value=f'  • {a}')
    c.font      = Font(name='Arial', size=10)
    c.fill      = make_fill('EBF3FB')
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws1.row_dimensions[i].height = 20

ws1.row_dimensions[3].height = 20
ws1.row_dimensions[4].height = 28
ws1.row_dimensions[5].height = 20

# Header row (row 6)
HDRS_S1 = ['Tahun', 'Total Klien Akumulatif\n(Pabrik)',
           'Total Mesin\nDipantau', 'Pendapatan Kotor (Rp)',
           'Beban Operasional / OpEx (Rp)', 'Aliran Kas Bersih / Proceeds (Rp)']
for col, h in enumerate(HDRS_S1, 1):
    hdr_cell(ws1, 6, col, h, bg=C_NAVY, sz=11)
ws1.row_dimensions[6].height = 40

# Data rows
s1_data = [
    # (tahun, klien, mesin, pendapatan, opex, proceeds, note_opex, note_proceeds)
    (0,  0,   0,    0,              2697900000,  -2697900000,  'Initial Inv.',   None),
    (1,  4,   20,   2400000000,     0,            2400000000,  'Sudah lunas di X0', None),
    (2,  8,   40,   4800000000,     2414610000,   2385390000,  'Naik 10%',       None),
    (3,  13,  65,   7800000000,     2656071000,   5143929000,  'Naik 10%',       None),
    (4,  20,  100,  12000000000,    2921678100,   9078321900,  'Naik 10%',       None),
    (5,  30,  150,  18000000000,    3213845910,   14786154090, 'Naik 10%',       None),
]

row_colors = [C_ORANGE, C_WHITE, C_GREEN, C_GREEN, C_GREEN, C_GREEN]
for i, (thn, klien, mesin, pend, opex, proc, note_o, _) in enumerate(s1_data):
    r  = 7 + i
    bg = row_colors[i]
    bld = (thn == 0)

    data_cell(ws1, r, 1, f'Tahun {thn}',       bg=bg, bold=bld, align='center')
    data_cell(ws1, r, 2, f'{klien} Pabrik',     bg=bg, align='center')
    data_cell(ws1, r, 3, f'{mesin} Mesin',      bg=bg, align='center')
    data_cell(ws1, r, 4, pend,                  bg=bg, align='right', fmt=NUM_FMT)
    # OpEx with note
    opex_label = f'{opex:,}  ({note_o})'.replace(',', '.') if note_o else f'{opex:,}'.replace(',', '.')
    c_opex = ws1.cell(row=r, column=5, value=opex if opex else None)
    c_opex.font      = Font(name='Arial', size=11)
    c_opex.fill      = make_fill(bg)
    c_opex.border    = border
    c_opex.alignment = Alignment(horizontal='right', vertical='center')
    c_opex.number_format = NUM_FMT

    data_cell(ws1, r, 6, proc, bg=bg, bold=bld, align='right', fmt=NUM_FMT)
    ws1.row_dimensions[r].height = 22

# Total row
tr = 7 + len(s1_data)
ws1.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=3)
c = ws1.cell(row=tr, column=1, value='TOTAL PROCEEDS (Tahun 1–5)')
c.font      = Font(name='Arial', bold=True, size=11)
c.fill      = make_fill(C_GOLD)
c.alignment = Alignment(horizontal='center', vertical='center')
c.border    = border
for col in [2, 3]:
    ws1.cell(row=tr, column=col).border = border

total_proc = sum(item[5] for item in s1_data if item[0] >= 1)
for col, val in zip([4, 5, 6], ['', '', total_proc]):
    c = data_cell(ws1, tr, col, val, bg=C_GOLD, bold=True, align='right', fmt=NUM_FMT)
ws1.row_dimensions[tr].height = 25

# Column widths S1
ws1.column_dimensions['A'].width = 12
ws1.column_dimensions['B'].width = 22
ws1.column_dimensions['C'].width = 16
ws1.column_dimensions['D'].width = 24
ws1.column_dimensions['E'].width = 28
ws1.column_dimensions['F'].width = 28
ws1.freeze_panes = 'A7'


# ════════════════════════════════════════════════════════════════════════════
#  SHEET 2 – PAYBACK PERIOD
# ════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('2. Payback Period')

title_row(ws2, 1, 'PAYBACK PERIOD (PP) – PT KAINOSOPH', 4)
sub_row(ws2, 2,
        'Tujuan: Menghitung berapa lama modal awal Rp 2.697.900.000 dapat dikembalikan.',
        4, bg='EBF3FB')

HDRS_S2 = ['No.', 'Keterangan', 'Aliran Kas (Rp)', 'Kumulatif Aliran Kas (Rp)']
for col, h in enumerate(HDRS_S2, 1):
    hdr_cell(ws2, 3, col, h, bg=C_NAVY, sz=11)
ws2.row_dimensions[3].height = 30

s2_data = [
    (1, 'Initial Investment (Tahun 0)', -2697900000,  -2697900000),
    (2, 'Proceeds Tahun ke-1',           2400000000,   -297900000),
    (3, 'Proceeds Tahun ke-2',           2385390000,  2087490000),
    (4, 'Proceeds Tahun ke-3',           5143929000,  7231419000),
    (5, 'Proceeds Tahun ke-4',           9078321900,  16309740900),
    (6, 'Proceeds Tahun ke-5',          14786154090,  31095894990),
]

for i, (no, ket, kas, kum) in enumerate(s2_data):
    r  = 4 + i
    bg = C_ORANGE if no == 1 else (C_GREEN if kum > 0 else C_WHITE)
    bld = (no == 1)
    data_cell(ws2, r, 1, no,   bg=bg, bold=bld, align='center')
    data_cell(ws2, r, 2, ket,  bg=bg, bold=bld, align='left')
    data_cell(ws2, r, 3, kas,  bg=bg, bold=bld, align='right', fmt=NUM_FMT)
    data_cell(ws2, r, 4, kum,  bg=bg, bold=bld, align='right', fmt=NUM_FMT)
    ws2.row_dimensions[r].height = 22

# Kesimpulan block
r_conc = 4 + len(s2_data) + 1
ws2.merge_cells(start_row=r_conc, start_column=1, end_row=r_conc, end_column=4)
c = ws2.cell(row=r_conc, column=1,
             value='Perhitungan Kekurangan: Defisit akhir Tahun 1 = Rp 297.900.000  |  '
                   'Kas masuk Tahun 2 = Rp 2.385.390.000  |  '
                   '(297.900.000 / 2.385.390.000) × 12 = 1,498 bulan ≈ 1 Bulan 15 Hari')
c.font      = Font(name='Arial', size=10, italic=True)
c.fill      = make_fill('FFF2CC')
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
c.border    = border
ws2.row_dimensions[r_conc].height = 35

r_pp = r_conc + 1
ws2.merge_cells(start_row=r_pp, start_column=1, end_row=r_pp, end_column=4)
c = ws2.cell(row=r_pp, column=1,
             value='✅  KESIMPULAN PP:  Modal Rp 2.697.900.000 kembali dalam  1 Tahun 1 Bulan 15 Hari  → SANGAT LAYAK')
c.font      = Font(name='Arial', bold=True, size=12, color='1F4E79')
c.fill      = make_fill(C_GOLD)
c.alignment = Alignment(horizontal='center', vertical='center')
c.border    = border
ws2.row_dimensions[r_pp].height = 30

ws2.column_dimensions['A'].width = 6
ws2.column_dimensions['B'].width = 30
ws2.column_dimensions['C'].width = 26
ws2.column_dimensions['D'].width = 28
ws2.freeze_panes = 'A4'


# ════════════════════════════════════════════════════════════════════════════
#  SHEET 3 – NPV, PI, IRR
# ════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet('3. NPV PI IRR')

title_row(ws3, 1, 'ANALISIS KELAYAKAN INVESTASI – NPV, PI & IRR  (Suku Bunga BI: 5,25%)', 5)
sub_row(ws3, 2, 'Tingkat Diskonto (r) = 5,25%  (0,0525)', 5, bg='EBF3FB')

HDRS_S3 = ['No.', 'Keterangan',
           'Aliran Kas (Rp)', 'Rumus Present Value (PV)', 'Present Value (Rp)']
for col, h in enumerate(HDRS_S3, 1):
    hdr_cell(ws3, 3, col, h, bg=C_NAVY, sz=11)
ws3.row_dimensions[3].height = 30

s3_data = [
    (1, 'Initial Investment (Tahun 0)',  -2697900000, 'X0 / (1+0,0525)^0',               -2697900000),
    (2, 'Proceeds Tahun ke-1',           2400000000, '2.400.000.000 / (1+0,0525)^1',      2280285035),
    (3, 'Proceeds Tahun ke-2',           2385390000, '2.385.390.000 / (1+0,0525)^2',      2153352542),
    (4, 'Proceeds Tahun ke-3',           5143929000, '5.143.929.000 / (1+0,0525)^3',      4411930298),
    (5, 'Proceeds Tahun ke-4',           9078321900, '9.078.321.900 / (1+0,0525)^4',      7398047886),
    (6, 'Proceeds Tahun ke-5',          14786154090, '14.786.154.090 / (1+0,0525)^5',    11448398336),
]

for i, (no, ket, kas, rumus, pv) in enumerate(s3_data):
    r  = 4 + i
    bg = C_ORANGE if no == 1 else C_WHITE
    bld = (no == 1)
    data_cell(ws3, r, 1, no,    bg=bg, bold=bld, align='center')
    data_cell(ws3, r, 2, ket,   bg=bg, bold=bld, align='left')
    data_cell(ws3, r, 3, kas,   bg=bg, bold=bld, align='right', fmt=NUM_FMT)
    data_cell(ws3, r, 4, rumus, bg=bg,           align='center', sz=10)
    data_cell(ws3, r, 5, pv,    bg=bg, bold=bld, align='right', fmt=NUM_FMT)
    ws3.row_dimensions[r].height = 22

# Total PV Proceeds row
r_tot = 4 + len(s3_data)
total_pv = 27692014097
for col in [1, 3, 4]:
    data_cell(ws3, r_tot, col, '', bg=C_LBLUE)
data_cell(ws3, r_tot, 2, 'TOTAL PV PROCEEDS', bg=C_LBLUE, bold=True, align='left')
data_cell(ws3, r_tot, 5, total_pv,             bg=C_LBLUE, bold=True, align='right', fmt=NUM_FMT)
ws3.row_dimensions[r_tot].height = 25

# ── Summary metrics ──────────────────────────────────────────────────────────
r_gap = r_tot + 1
ws3.row_dimensions[r_gap].height = 10          # spacer

metrics = [
    ('NPV',
     'Net Present Value',
     '= Total PV Proceeds − Initial Investment\n= Rp 27.692.014.097 − Rp 2.697.900.000',
     24994114097,
     '✅ SANGAT LAYAK  (NPV > 0)',
     C_GREEN),
    ('PI',
     'Profitability Index',
     '= Total PV Proceeds / Initial Investment\n= 27.692.014.097 / 2.697.900.000',
     10.26,
     '✅ SANGAT LAYAK  (PI > 1)',
     C_GREEN),
    ('IRR',
     'Internal Rate of Return',
     'Tingkat diskonto yang membuat NPV = 0\n(dihitung via regresi finansial atas arus kas Tahun 0–5)',
     1.2148,
     '✅ SANGAT LAYAK  (IRR 121,48% >> r 5,25%)',
     C_GREEN),
]

# Sub-header
r_mhdr = r_gap + 1
hdr_cell(ws3, r_mhdr, 1, 'Metrik',     bg=C_BLUE, sz=11)
hdr_cell(ws3, r_mhdr, 2, 'Nama Lengkap', bg=C_BLUE, sz=11)
hdr_cell(ws3, r_mhdr, 3, 'Formula / Cara Hitung', bg=C_BLUE, sz=11)
hdr_cell(ws3, r_mhdr, 4, 'Nilai',      bg=C_BLUE, sz=11)
hdr_cell(ws3, r_mhdr, 5, 'Kesimpulan', bg=C_BLUE, sz=11)
ws3.row_dimensions[r_mhdr].height = 28

for j, (kode, nama, formula, nilai, konklusi, bg) in enumerate(metrics):
    r  = r_mhdr + 1 + j
    data_cell(ws3, r, 1, kode,    bg=bg, bold=True, align='center', sz=12)
    data_cell(ws3, r, 2, nama,    bg=bg, bold=True, align='left')
    data_cell(ws3, r, 3, formula, bg=bg,             align='left', sz=10)

    # Format nilai
    if kode == 'NPV':
        data_cell(ws3, r, 4, nilai, bg=bg, bold=True, align='right', fmt=NUM_FMT)
    elif kode == 'PI':
        data_cell(ws3, r, 4, nilai, bg=bg, bold=True, align='right', fmt='0.00')
    else:  # IRR
        data_cell(ws3, r, 4, nilai, bg=bg, bold=True, align='right', fmt='0.00%')

    data_cell(ws3, r, 5, konklusi, bg=bg, bold=True, align='center')
    ws3.row_dimensions[r].height = 42

ws3.column_dimensions['A'].width = 8
ws3.column_dimensions['B'].width = 26
ws3.column_dimensions['C'].width = 40
ws3.column_dimensions['D'].width = 24
ws3.column_dimensions['E'].width = 36
ws3.freeze_panes = 'A4'


# ── Save ─────────────────────────────────────────────────────────────────────
OUT = (r'C:\PORTFOLIO\PROJECTS\PBL\Predictive Maintenance'
       r'\projects\predictive-maintenance-monorepo\kewirus\InvestKainosoph.xlsx')
wb.save(OUT)
print('SAVED:', OUT)
