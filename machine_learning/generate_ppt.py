import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx is not installed. Please install it to continue.")
    sys.exit(1)

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*tuple(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

def create_ppt(md_filepath, output_filepath):
    prs = Presentation()
    primary_color = hex_to_rgb('265528')
    
    with open(md_filepath, 'r', encoding='utf-8') as f:
        content = f.read()
    
    slides_raw = content.split('### Slide ')
    
    for slide_raw in slides_raw[1:]: 
        lines = slide_raw.strip().split('\n')
        if not lines: continue
        
        # Parse title
        title_line = lines[0].strip()
        if '—' in title_line:
            title_text = title_line.split('—', 1)[1].strip()
        elif '-' in title_line:
            title_text = title_line.split('-', 1)[1].strip()
        else:
            title_text = title_line
            
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(255, 255, 255)
        
        # Title
        title_shape = slide.shapes.title
        title_shape.text = title_text
        for run in title_shape.text_frame.paragraphs[0].runs:
            run.font.color.rgb = primary_color
            run.font.bold = True
            run.font.name = 'Arial'
            
        # Body Content
        body_shape = slide.shapes.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        
        speaker_notes = []
        vis_text = ""
        
        for line in lines[1:]:
            line = line.strip()
            # Bullet point logic
            if line.startswith('- ') or line.startswith('* '):
                if '[VIS]' in line:
                    vis_text = line.replace('- ', '', 1).replace('* ', '', 1).replace('[VIS]', '').strip()
                else:
                    p = tf.add_paragraph()
                    p.text = line.replace('- ', '', 1).replace('* ', '', 1).replace('**', '')
                    p.font.color.rgb = primary_color
                    p.font.size = Pt(16)
                    p.level = 0
            # Indented bullets
            elif line.startswith('  - ') or line.startswith('    - '):
                p = tf.add_paragraph()
                p.text = line.strip().replace('- ', '', 1).replace('**', '')
                p.font.color.rgb = primary_color
                p.font.size = Pt(14)
                p.level = 1
            # Speaker notes
            elif line.startswith('> **Speaker Notes:**') or line.startswith('> Speaker Notes:'):
                clean_note = line.replace('> **Speaker Notes:**', '').replace('> Speaker Notes:', '').strip()
                if clean_note:
                    speaker_notes.append(clean_note)
            elif line.startswith('>'):
                clean_note = line.replace('>', '').strip()
                if clean_note:
                    speaker_notes.append(clean_note)
                
        # Handle VIS block
        if vis_text:
            left = Inches(1)
            top = Inches(4.5)
            width = Inches(8)
            height = Inches(2.5)
            try:
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
                shape.fill.background()
                shape.line.color.rgb = primary_color
                shape.line.width = Pt(2)
                
                text_frame = shape.text_frame
                text_frame.text = f"--- FRAME VISUAL KOSONG ---\n({vis_text})"
                for paragraph in text_frame.paragraphs:
                    paragraph.alignment = PP_ALIGN.CENTER
                    for run in paragraph.runs:
                        run.font.color.rgb = primary_color
                        run.font.size = Pt(14)
                        run.font.bold = True
            except Exception as e:
                print(f"Warning: Could not add shape for VIS: {e}")
                
        # Handle Speaker Notes
        if speaker_notes:
            try:
                notes_slide = slide.notes_slide
                notes_text_frame = notes_slide.notes_text_frame
                notes_text_frame.text = " ".join(speaker_notes)
            except Exception as e:
                print(f"Warning: Could not add speaker notes: {e}")

    prs.save(output_filepath)
    print(f"\nBerhasil membuat presentasi: {output_filepath}")

if __name__ == '__main__':
    base_dir = r"c:\PORTFOLIO\PROJECTS\PBL\Predictive Maintenance\projects\predictive-maintenance-monorepo\machine_learning"
    md_file = os.path.join(base_dir, "PPT_CRISP_DM_OUTLINE.md")
    out_file = os.path.join(base_dir, "Presentasi_CRISP_DM_LapisAI.pptx")
    
    if os.path.exists(md_file):
        create_ppt(md_file, out_file)
    else:
        print(f"Error: File {md_file} not found.")
